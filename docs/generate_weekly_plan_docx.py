import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_weekly_plan_document():
    doc = docx.Document()

    # Academic margins: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(6)

    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    BLUE = RGBColor(0x2E, 0x75, 0xB6)
    DARK = RGBColor(0x22, 0x22, 0x22)
    GRAY = RGBColor(0x55, 0x55, 0x55)

    def set_cell_background(cell, fill_hex):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, color="D3D3D3"):
        tblPr = table._tbl.tblPr
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
        '''
        tblPr.append(parse_xml(borders_xml))

    # Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst1 = p_inst.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN – ĐHQG-HCM\n")
    r_inst1.font.size = Pt(11)
    r_inst1.font.bold = True
    r_inst2 = p_inst.add_run("HỌC PHẦN: MÁY HỌC NÂNG CAO (ADVANCED MACHINE LEARNING)\n")
    r_inst2.font.size = Pt(11)
    r_inst2.font.bold = True
    r_inst2.font.color.rgb = BLUE
    p_inst.paragraph_format.space_after = Pt(14)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("KẾ HOẠCH TRIỂN KHAI STEP-BY-STEP CHI TIẾT THEO TUẦN VÀ TỪNG THÀNH VIÊN\n")
    r_t1.font.size = Pt(13)
    r_t1.font.bold = True
    r_t1.font.color.rgb = NAVY
    
    r_t2 = p_title.add_run("ĐỒ ÁN: KHUNG HỌC MÁY ĐA PHƯƠNG THỨC NHẬN DIỆN TRẠNG THÁI RỦI RO AN TOÀN CHO TÀI XẾ XE HAI BÁNH\n")
    r_t2.font.size = Pt(14)
    r_t2.font.bold = True
    r_t2.font.color.rgb = NAVY

    r_t3 = p_title.add_run("Phân công chi tiết 8 tuần (Tuần 2 đến Tuần 9) cho 5 thành viên nhóm Trí tuệ Nhân tạo")
    r_t3.font.size = Pt(10.5)
    r_t3.font.italic = True
    r_t3.font.color.rgb = GRAY
    p_title.paragraph_format.space_after = Pt(14)

    # Info table
    info_table = doc.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Giảng viên hướng dẫn:", "ThS./TS. Đặng Việt Dũng"),
        ("Khung thời gian thực hiện:", "27/10/2026 – 20/12/2026 (8 Tuần làm việc)"),
        ("Nhóm sinh viên thực hiện:", "Nhóm 5 thành viên (Trường ĐH Công nghệ Thông tin – ĐHQG-HCM)"),
        ("Trưởng nhóm điều phối:", "Trần Hoàng Hôn – MSSV: 26410046"),
        ("Các thành viên chuyên trách:", "Nguyễn Duy Khang (26410055) – Computer Vision Lead\nPhạm Thành Trung (26410141) – Bio-Signals & Time-Series\nVũ Văn Duy (26410031) – ML Modeling & Evaluation\nHoàng Cao Sơn (25730061) – Data Engineering & XAI"),
        ("Mục tiêu nghiệm thu:", "Báo cáo Word/PDF >= 40 trang, Codebase GitHub đạt Macro-F1 >= 90%")
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.3)
        c1.width = Inches(4.5)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        set_cell_background(c0, "F2F5F9")
        set_cell_background(c1, "FAFAFA")
        set_cell_margins(c0, 45, 45, 90, 90)
        set_cell_margins(c1, 45, 45, 90, 90)
    set_table_borders(info_table, "B0C4DE")
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_h1(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    def add_h2(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = BLUE
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)

    def add_p(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r2 = p.add_run(text)
        r2.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

    # 1. Mục tiêu
    add_h1("1. MỤC TIÊU VÀ PHÂN CÔNG CHUYÊN TRÁCH 5 THÀNH VIÊN")
    add_p(
        "Kế hoạch hành động chi tiết này nhằm đảm bảo 5 thành viên làm việc nhịp nhàng, có sản phẩm bàn giao cụ thể "
        "theo từng tuần, không trùng lặp đầu việc và bảo đảm tiến độ hoàn thành đồ án học máy chất lượng cao."
    )
    add_bullet("Trần Hoàng Hôn (Trưởng nhóm / ML Architect): ", "Kiến trúc tổng thể, mô hình hóa toán học, khối Cross-Modal Attention, logic ngắt SOS, chủ biên báo cáo và điều phối.")
    add_bullet("Nguyễn Duy Khang (Computer Vision Lead): ", "Pipeline dữ liệu ảnh nón bảo hiểm và hung khí, Spatial Augmentation chống rung lắc, MobileNetV3-Small backbone và Vision Baseline.")
    add_bullet("Phạm Thành Trung (Bio-Signals Specialist): ", "Pipeline chuỗi thời gian PPG nhịp tim, lọc Butterworth, trích xuất 8 đặc trưng HRV (SDNN, RMSSD, pNN50), huấn luyện BiLSTM.")
    add_bullet("Vũ Văn Duy (ML Modeling & Evaluation): ", "Cài đặt Focal Loss, Class-Balanced Loss, Cost-Sensitive Bayes Risk Matrix, kiểm định Stratified 5-Fold, PR-AUC, Confusion Matrix.")
    add_bullet("Hoàng Cao Sơn (Data Engineering & XAI): ", "Data Pipeline, trực quan hóa Grad-CAM (ảnh) và SHAP (nhịp tim), thực nghiệm hiện trường TP.HCM, dựng video demo và Release.")

    # 2. Chi tiết từng tuần
    add_h1("2. KẾ HOẠCH STEP-BY-STEP THEO TỪNG TUẦN (TUẦN 2 ĐẾN TUẦN 9)")

    weeks_data = [
        ("TUẦN 2 (27/10 – 01/11/2026): KHỞI ĐỘNG, CHỐT BÀI TOÁN & TẢI DỮ LIỆU GỐC",
         "Thiết lập Git repo, cài đặt môi trường ảo Python/PyTorch, tải về các bộ dữ liệu thô (NTHU-DDD, Roboflow Weapon, WESAD).",
         [("Trần Hoàng Hôn", "Thiết lập Git workflow, quy chuẩn commit, rà soát thuyết minh toán học và chủ trì họp nhóm.", "CONTRIBUTING.md, README.md", "Repo GitHub chuẩn hóa"),
          ("Nguyễn Duy Khang", "Tải bộ dữ liệu NTHU-DDD (tài xế đeo kính/khẩu trang) và Roboflow Weapon (>4.200 ảnh).", "src/data/download_vision.py", "Dữ liệu ảnh thô data/raw/vision/"),
          ("Phạm Thành Trung", "Tải bộ dữ liệu WESAD, viết script đọc file .pkl của 15 đối tượng, trích xuất chuỗi nhịp tim PPG.", "src/data/parse_wesad.py", "File CSV nhịp tim thô data/raw/bio/"),
          ("Vũ Văn Duy", "Tạo requirements.txt, environment.yml, viết test_env.py kiểm tra CUDA và thư viện PyTorch/Scikit-learn.", "requirements.txt, test_env.py", "Môi trường test pass 100%"),
          ("Hoàng Cao Sơn", "Khởi tạo notebook EDA 01, lên kế hoạch thu thập ảnh nón bảo hiểm thực tế ngoài đường TP.HCM.", "notebooks/01_eda.ipynb", "Khung notebook EDA & 50 ảnh mẫu")],
         "Toàn bộ 5 thành viên clone repo về máy, chạy test_env.py pass 100%, dữ liệu thô đã nằm trong data/raw/."),

        ("TUẦN 3 (02/11 – 08/11/2026): TIỀN XỬ LÝ DỮ LIỆU, DATA AUGMENTATION & CHIA 5-FOLD",
         "Chuẩn hóa ảnh về 96x96, cắt chuỗi nhịp tim cửa sổ trượt T=30s, cài đặt Augmentation chống rung xe, chia Stratified 5-Fold.",
         [("Trần Hoàng Hôn", "Định nghĩa Data Interface chuẩn {'image', 'bio_seq', 'sos', 'label'}, chia StratifiedGroupKFold (K=5).", "src/data/splitter.py", "File splits_5fold.json (no data leakage)"),
          ("Nguyễn Duy Khang", "Crop vùng nón bảo hiểm, resize 3x96x96, cài đặt Motion Blur (5x5, 9x9), ánh sáng ±40%, shadow injection.", "src/data/vision_transforms.py", "Bộ ảnh tiền xử lý data/processed/images/"),
          ("Phạm Thành Trung", "Lọc thông dải Butterworth (0.5-4.0 Hz) trên PPG, cắt cửa sổ trượt T=30s, thêm Gaussian Noise và time-warping.", "src/data/bio_transforms.py", "Chuỗi nhịp tim sạch data/processed/bio/"),
          ("Vũ Văn Duy", "Cài đặt PyTorch Dataset và DataLoader nạp đồng thời ảnh và chuỗi nhịp tim, kiểm tra batch tensor shape.", "src/data/multimodal_dataset.py", "Unit test DataLoader pass 100%"),
          ("Hoàng Cao Sơn", "Hoàn thiện notebook 01_eda, vẽ biểu đồ phân bố nhãn 70:20:10, trực quan hóa ảnh trước/sau Augmentation.", "notebooks/01_eda.ipynb", "Biểu đồ EDA trong reports/figures/")],
         "Chạy pytest tests/test_dataset.py pass; file splits_5fold.json hoàn chỉnh; notebook EDA trực quan hóa đầy đủ."),

        ("TUẦN 4 (09/11 – 15/11/2026): FEATURE ENGINEERING & HUẤN LUYỆN MÔ HÌNH CƠ SỞ (BASELINES)",
         "Trích xuất 8 đặc trưng HRV, đặc trưng HOG/LBP, huấn luyện LightGBM/XGBoost và Single-modal MobileNetV3 làm benchmark.",
         [("Trần Hoàng Hôn", "Xây dựng hàm tính toán Macro-F1, Recall từng lớp, Precision từng lớp; thiết lập bảng theo dõi benchmark.", "src/evaluation/metrics.py", "Khuôn khổ đo lường chuẩn hóa"),
          ("Nguyễn Duy Khang", "Trích xuất HOG & LBP trên nón bảo hiểm; huấn luyện mô hình MobileNetV3-Small (Vision-only) trên 5-Fold.", "src/models/vision_only.py", "Checkpoint models/vision_only.pth"),
          ("Phạm Thành Trung", "Tích hợp compute_hrv_features trích xuất 8 chỉ số HRV; huấn luyện BiLSTM (Bio-only) trên chuỗi nhịp tim.", "src/models/bio_only.py", "Checkpoint models/bio_only.pth"),
          ("Vũ Văn Duy", "Ghép vector đặc trưng thủ công (HOG + LBP + HRV); huấn luyện LightGBM và XGBoost trên dữ liệu bảng.", "notebooks/02_baseline_models.ipynb", "Mô hình LightGBM/XGBoost được train"),
          ("Hoàng Cao Sơn", "Vẽ ma trận nhầm lẫn Confusion Matrix cho từng baseline, phân tích các trường hợp báo động sai của mô hình đơn lẻ.", "reports/figures/", "Phân tích nhược điểm mô hình đơn lẻ")],
         "Hoàn thành notebook 02_baseline_models.ipynb; ghi nhận F1-Score của 4 mô hình cơ sở để làm mốc đối chứng."),

        ("TUẦN 5 (16/11 – 22/11/2026): PHÁT TRIỂN KIẾN TRÚC HỌC SÂU ĐA PHƯƠNG THỨC & CUSTOM LOSSES",
         "Cài đặt 3 cấu trúc dung hợp (Early, Late, Cross-Modal Attention), tích hợp Focal Loss, Class-Balanced Loss và Ma trận Chi phí Bayes.",
         [("Trần Hoàng Hôn", "Cài đặt Cross-Modal Attention (Bio Query -> Visual Key/Value), nhánh ngắt nút SOS, lớp MultimodalDriverSafetyNet.", "src/models/multimodal_net.py", "Mô hình Cross-Attention hoàn chỉnh"),
          ("Nguyễn Duy Khang", "Xây dựng VisualFeatureExtractor sử dụng MobileNetV3-Small và ResNet-18, tính toán số tham số và FLOPs.", "src/models/backbones/", "Backbone thị giác tối ưu"),
          ("Phạm Thành Trung", "Xây dựng TemporalBioExtractor (1D-CNN + BiLSTM 32 units), cài đặt cơ chế Early Fusion (nối phẳng).", "src/models/temporal_extractor.py", "Module trích xuất chuỗi nhịp tim"),
          ("Vũ Văn Duy", "Tối ưu hóa MultiClassFocalLoss, ClassBalancedFocalLoss và CostSensitiveBayesLoss (phạt nặng FN lớp nguy hiểm C_20=200).", "src/models/losses.py", "Trainer module hỗ trợ AMP fp16"),
          ("Hoàng Cao Sơn", "Cài đặt mô hình Late Fusion (Soft-voting thích nghi), viết unit test forward pass cho cả 3 mô hình dung hợp.", "tests/test_models.py", "Unit tests forward pass pass 100%")],
         "Cả 3 cấu trúc dung hợp chạy mượt trên dữ liệu thử nghiệm; Trainer chạy được 1 epoch không lỗi tensor hay NaN loss."),

        ("TUẦN 6 (23/11 – 29/11/2026): HUẤN LUYỆN END-TO-END 5-FOLD & TINH CHỈNH SIÊU THAM SỐ",
         "Huấn luyện toàn diện mô hình Cross-Attention trên 5 Folds độc lập, tinh chỉnh siêu tham số, lưu checkpoints tốt nhất.",
         [("Trần Hoàng Hôn", "Điều phối huấn luyện 5 Folds trên GPU, theo dõi Learning Curve, lưu checkpoint Macro-F1 cao nhất.", "notebooks/03_multimodal.ipynb", "Trọng số 5 Folds models/checkpoints/"),
          ("Nguyễn Duy Khang", "Fine-tuning các tầng sâu MobileNetV3 với learning rate nhỏ 1e-4, kiểm tra độ nhạy với ảnh mờ/ngược sáng.", "src/training/finetune_vision.py", "Backbone fine-tuned ổn định"),
          ("Phạm Thành Trung", "Thực nghiệm so sánh kích thước cửa sổ trượt T (15s vs 30s vs 60s), xác nhận T=30s cho cân bằng tốt nhất.", "notebooks/experiments_hr.ipynb", "Báo cáo thực nghiệm cửa sổ trượt T"),
          ("Vũ Văn Duy", "Thực hiện Ablation Study so sánh: Cross-Entropy vs Focal Loss (gamma=2.0) vs Class-Balanced Loss.", "reports/figures/loss_ablation.png", "Báo cáo Ablation Study chứng minh Focal Loss"),
          ("Hoàng Cao Sơn", "Đo lường tốc độ suy luận (FPS & Latency ms) trên CPU mô phỏng thiết bị biên, kiểm tra RAM sử dụng.", "src/evaluation/benchmark_speed.py", "Đạt >= 15 FPS, latency < 70 ms trên CPU")],
         "Hoàn thành huấn luyện 5 Folds; mô hình Cross-Attention đạt Macro-F1 >= 90%, Recall Drowsy >= 95%, Recall Assault >= 98%."),

        ("TUẦN 7 (30/11 – 06/12/2026): THẨM ĐỊNH CHUYÊN SÂU, THỰC NGHIỆM ĐƯỜNG PHỐ & XAI",
         "Đánh giá trên Held-out Test Set, vẽ PR-AUC curve, trực quan hóa Grad-CAM & SHAP, kiểm thử dữ liệu thực tế tại TP.HCM.",
         [("Trần Hoàng Hôn", "Hoàn thiện notebook 04_evaluation, tính trung bình và độ lệch chuẩn qua 5 Folds (Mean ± Std), tính chi phí Bayes.", "notebooks/04_evaluation.ipynb", "Bảng tổng hợp số liệu kiểm định cuối"),
          ("Nguyễn Duy Khang", "Cài đặt Grad-CAM trên MobileNetV3, xuất ảnh heatmap chứng minh mô hình soi đúng nón bảo hiểm và hung khí.", "src/evaluation/gradcam.py", "Ảnh heatmap Grad-CAM trong reports/figures/"),
          ("Phạm Thành Trung", "Cài đặt SHAP values trên đặc trưng nhịp tim, xuất biểu đồ summary/waterfall giải thích kìm hãm báo động giả.", "src/evaluation/shap_analysis.py", "Biểu đồ SHAP trong reports/figures/"),
          ("Vũ Văn Duy", "Vẽ đường cong PR-AUC cho 3 lớp, vẽ Normalized Confusion Matrix trực quan hóa bằng seaborn heatmap.", "reports/figures/pr_curve.png", "Đồ thị PR-AUC & Ma trận nhầm lẫn chuẩn"),
          ("Hoàng Cao Sơn", "Chạy kiểm thử mô hình trên tập video nón bảo hiểm thực tế thu thập tại TP.HCM (OOD Testing), ghi nhận kết quả.", "src/evaluation/test_realworld.py", "Báo cáo thực nghiệm đường phố TP.HCM")],
         "Hoàn thành notebook 04; bộ ảnh Grad-CAM, biểu đồ SHAP, đồ thị PR-AUC sẵn sàng đưa vào báo cáo và slide."),

        ("TUẦN 8 (07/12 – 13/12/2026): HOÀN THIỆN QUYỂN BÁO CÁO KHOA HỌC (>= 40 TRANG) & ĐÓNG GÓI MÃ NGUỒN",
         "Biên tập quyển báo cáo khoa học Word + PDF >= 40 trang chuẩn UIT, đóng gói mã nguồn và tài liệu tái lập kết quả.",
         [("Trần Hoàng Hôn", "Chủ biên báo cáo: Viết Mở đầu, Tóm tắt, Chương 1, Chương 3 (Toán học), Chương 5 (Kiến trúc) và Kết luận.", "docs/Bao_Cao_Do_An_ML.docx", "Bản thảo báo cáo khoa học >= 40 trang"),
          ("Nguyễn Duy Khang", "Soạn thảo Chương 4 (Dữ liệu thị giác): Mô tả NTHU-DDD, Roboflow Weapon, Spatial Augmentation kèm ảnh mẫu.", "docs/chapters/chap4_vision.md", "Nội dung & hình ảnh Chương 4"),
          ("Phạm Thành Trung", "Soạn thảo Chương 4 (Dữ liệu sinh học): Mô tả WESAD, lọc Butterworth, công thức 8 đặc trưng HRV kèm đồ thị sóng.", "docs/chapters/chap4_bio.md", "Nội dung & đồ thị nhánh nhịp tim"),
          ("Vũ Văn Duy", "Soạn thảo Chương 6 (Thực nghiệm & Đánh giá): Bảng kết quả 5 Folds, PR-AUC, Confusion Matrix, Ablation Study.", "docs/chapters/chap6_eval.md", "Bảng biểu số liệu và phân tích kết quả"),
          ("Hoàng Cao Sơn", "Soạn thảo Chương 7 (XAI & Ứng dụng thực tế): Phân tích Grad-CAM, SHAP, kết quả TP.HCM, chuẩn hóa tài liệu tham khảo.", "docs/chapters/chap7_xai.md", "Nội dung XAI & danh mục trích dẫn")],
         "Xuất bản file Bao_Cao_Do_An_ML.docx và PDF >= 40 trang, đầy đủ bảng biểu số liệu thực nghiệm và đồ thị minh chứng."),

        ("TUẦN 9 (14/12 – 20/12/2026): SLIDE BÁO CÁO, VIDEO DEMO & BẢO VỆ ĐỒ ÁN CUỐI KỲ",
         "Làm slide thuyết trình chuyên nghiệp, quay video demo mô hình thực tế, diễn tập bảo vệ và bảo vệ trước Giảng viên.",
         [("Trần Hoàng Hôn", "Thiết kế bộ Slide thuyết trình chính (25-30 slides), phân chia lượt nói cho 5 bạn, chủ trì buổi bảo vệ thử.", "docs/Slide_Bao_Ve_ML.pptx", "Slide PowerPoint chuyên nghiệp"),
          ("Nguyễn Duy Khang", "Chuẩn bị phần thuyết trình nhánh Thị giác và câu hỏi phản biện liên quan đến góc nghiêng nón bảo hiểm.", "docs/qa_prep.md", "Nội dung thuyết trình nhánh Vision"),
          ("Phạm Thành Trung", "Chuẩn bị phần thuyết trình nhánh Tín hiệu sinh học và câu hỏi phản biện về độ trễ nhịp tim.", "docs/qa_prep.md", "Nội dung thuyết trình nhánh Bio-signals"),
          ("Vũ Văn Duy", "Chuẩn bị phần thuyết trình Mô hình toán học, Focal Loss, Ma trận chi phí Bayes và giải thích bảng số liệu.", "docs/qa_prep.md", "Nội dung thuyết trình nhánh ML Modeling"),
          ("Hoàng Cao Sơn", "Quay và dựng Video clip demo thực tế (2-3 phút), đóng gói mã nguồn Release v1.0.0 trên GitHub.", "docs/demo_video.mp4, GitHub v1.0.0", "Video clip demo & GitHub Release v1.0.0")],
         "Bảo vệ thử đạt thời gian 15 phút trình bày + 10 phút Q&A; nộp đầy đủ Báo cáo PDF, Slide, Source code và Video cho Thầy Dũng.")
    ]

    for title, desc, tasks, dod in weeks_data:
        add_h2(title)
        add_p(f"Mục tiêu tuần: {desc}")
        
        tbl = doc.add_table(rows=len(tasks)+1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Thành viên", "Nhiệm vụ cụ thể (Action Items)", "File mã nguồn phụ trách", "Sản phẩm bàn giao (Deliverables)"]
        for j, h in enumerate(headers):
            cell = tbl.cell(0, j)
            p = cell.paragraphs[0]
            r = p.add_run(h)
            r.font.bold = True
            r.font.size = Pt(9)
            set_cell_background(cell, "1B365D")
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_margins(cell, 50, 50, 70, 70)
        
        for i, row_data in enumerate(tasks):
            row = tbl.rows[i+1]
            for j, val in enumerate(row_data):
                cell = row.cells[j]
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if j == 0:
                    r.font.bold = True
                    set_cell_background(cell, "F2F5F9")
                elif j == 2:
                    r.font.italic = True
                    set_cell_background(cell, "FAFAFA")
                else:
                    set_cell_background(cell, "FAFAFA")
                set_cell_margins(cell, 40, 40, 70, 70)
        set_table_borders(tbl, "B0C4DE")
        
        p_dod = doc.add_paragraph()
        r_dod_t = p_dod.add_run("Tiêu chí hoàn thành (Definition of Done - DoD): ")
        r_dod_t.font.bold = True
        r_dod_t.font.size = Pt(9.5)
        r_dod_t.font.color.rgb = BLUE
        r_dod = p_dod.add_run(dod)
        r_dod.font.size = Pt(9.5)
        r_dod.font.italic = True
        p_dod.paragraph_format.space_before = Pt(3)
        p_dod.paragraph_format.space_after = Pt(8)

    # 3. Quy chế phối hợp
    add_h1("3. QUY CHẾ PHỐI HỢP & NGUYÊN TẮC LÀM VIỆC NHÓM")
    add_bullet("Họp giao ban định kỳ: ", "Tối Chủ Nhật 20h00 đánh giá tuần cũ và giao việc tuần mới; Tối Thứ Tư 21h00 đồng bộ 5 phút giải quyết blocker.")
    add_bullet("Quy chuẩn Git: ", "Không push trực tiếp vào main. Tạo nhánh feature/<ten-nhanh>, tạo Pull Request với tối thiểu 1 review và Trưởng nhóm duyệt merge.")
    add_bullet("Đo lường minh bạch: ", "Đánh giá chéo Peer Review dựa trên đúng hạn (On-time) và chất lượng code/nội dung đóng góp.")

    docs_output = r"D:\Trí tuệ nhân tạo\HK3\Máy học nâng cao\Đồ án\docs\Ke_Hoach_Chi_Tiet_Tung_Tuan_Va_Thanh_Vien.docx"
    root_output = r"D:\Trí tuệ nhân tạo\HK3\Máy học nâng cao\Đồ án\Ke_Hoach_Chi_Tiet_Tung_Tuan_Va_Thanh_Vien.docx"
    doc.save(docs_output)
    doc.save(root_output)
    print("Weekly plan docx created successfully.")

if __name__ == "__main__":
    create_weekly_plan_document()
