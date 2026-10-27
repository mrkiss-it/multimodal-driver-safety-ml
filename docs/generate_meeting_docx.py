import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_meeting_document():
    doc = docx.Document()

    # Academic margins
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(5)

    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    BLUE = RGBColor(0x2E, 0x75, 0xB6)
    DARK = RGBColor(0x22, 0x22, 0x22)
    GRAY = RGBColor(0x55, 0x55, 0x55)

    def set_cell_background(cell, fill_hex):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, color="D3D3D3"):
        tblPr = table._tbl.tblPr
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
        '''
        tblPr.append(parse_xml(borders_xml))

    # Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_header.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r1.font.bold = True
    r1.font.size = Pt(11)
    r2 = p_header.add_run("Độc lập – Tự do – Hạnh phúc\n")
    r2.font.bold = True
    r2.font.size = Pt(11)
    r3 = p_header.add_run("----------o0o----------\n\n")
    r3.font.size = Pt(10)
    p_header.paragraph_format.space_after = Pt(8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("BIÊN BẢN HỌP KHỞI ĐỘNG ĐỒ ÁN MÔN HỌC (TUẦN 2)\n")
    r_t1.font.bold = True
    r_t1.font.size = Pt(13)
    r_t1.font.color.rgb = NAVY
    
    r_t2 = p_title.add_run("HỌC PHẦN: MÁY HỌC NÂNG CAO (ADVANCED MACHINE LEARNING)\n")
    r_t2.font.bold = True
    r_t2.font.size = Pt(11.5)
    r_t2.font.color.rgb = BLUE

    r_t3 = p_title.add_run("Đề tài: Khung mô hình học máy đa phương thức nhận diện trạng thái rủi ro an toàn cho tài xế xe hai bánh")
    r_t3.font.italic = True
    r_t3.font.size = Pt(10)
    r_t3.font.color.rgb = GRAY
    p_title.paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    def add_h2(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = BLUE
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)

    def add_p(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r2 = p.add_run(text)
        r2.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

    # I. THÔNG TIN CHUNG
    add_h1("I. THÔNG TIN CHUNG VỀ CUỘC HỌP")
    add_bullet("Thời gian: ", "20h00 – 21h45, Thứ Ba ngày 27/10/2026.")
    add_bullet("Hình thức: ", "Họp trực tuyến qua Microsoft Teams / Google Meet.")
    add_bullet("Chủ tọa cuộc họp: ", "Trần Hoàng Hôn (Trưởng nhóm – MSSV: 26410046).")
    add_bullet("Thư ký ghi biên bản: ", "Hoàng Cao Sơn (Thành viên – MSSV: 25730061).")
    add_bullet("Thành phần tham dự: ", "Có mặt đầy đủ 05/05 thành viên nhóm.")

    info_table = doc.add_table(rows=6, cols=5)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["STT", "Họ và Tên", "MSSV", "Vai trò trong nhóm", "Tham dự"]
    for j, h in enumerate(headers):
        cell = info_table.cell(0, j)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        set_cell_background(cell, "1B365D")
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_margins(cell, 40, 40, 60, 60)

    members = [
        ("1", "Trần Hoàng Hôn", "26410046", "Trưởng nhóm / ML Architect", "Có mặt (Chủ tọa)"),
        ("2", "Nguyễn Duy Khang", "26410055", "Thành viên / Computer Vision Lead", "Có mặt"),
        ("3", "Phạm Thành Trung", "26410141", "Thành viên / Bio-Signal Specialist", "Có mặt"),
        ("4", "Vũ Văn Duy", "26410031", "Thành viên / ML Evaluation Lead", "Có mặt"),
        ("5", "Hoàng Cao Sơn", "25730061", "Thành viên / Data Engineering & XAI", "Có mặt (Thư ký)")
    ]
    for i, m in enumerate(members):
        row = info_table.rows[i+1]
        for j, val in enumerate(m):
            cell = row.cells[j]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 0:
                set_cell_background(cell, "F2F5F9")
            elif j == 1:
                r.font.bold = True
                set_cell_background(cell, "FAFAFA")
            else:
                set_cell_background(cell, "FAFAFA")
            set_cell_margins(cell, 35, 35, 60, 60)
    set_table_borders(info_table, "B0C4DE")
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # II. NỘI DUNG CUỘC HỌP
    add_h1("II. NỘI DUNG CUỘC HỌP & CÁC QUYẾT NGHỊ CHÍNH")
    
    add_h2("1. Thống nhất Tên đề tài và Mục tiêu Học thuật")
    add_p(
        "Trưởng nhóm Trần Hoàng Hôn trình bày định hướng đề tài giải quyết vấn đề thực tế nóng hổi: "
        "bảo vệ an toàn kép cho tài xế xe hai bánh tại Việt Nam (ngủ gật tích tắc và nguy cơ cướp giật ban đêm). "
        "Đề tài hoàn toàn tránh các hướng làm sáo rỗng và đáp ứng tiêu chí đánh giá khắt khe của Thầy Đặng Việt Dũng."
    )
    add_bullet("Tên đề tài tiếng Việt: ", "Khung mô hình học máy đa phương thức nhận diện trạng thái rủi ro an toàn cho tài xế xe hai bánh dựa trên thị giác máy tính và chuỗi thời gian sinh trắc học.")
    add_bullet("Tên đề tài tiếng Anh: ", "A Multimodal Machine Learning Framework for Two-Wheel Driver Safety Risk State Recognition Integrating Computer Vision and Physiological Bio-Signal Time-Series.")
    add_bullet("Mục tiêu điểm số kỳ vọng: ", "9.5 – 10.0 điểm.")

    add_h2("2. Thông qua Phân công Vai trò Chuyên môn 5 Thành viên")
    add_bullet("Trần Hoàng Hôn: ", "ML Architect, quản lý tiến độ Git, thiết kế khối Cross-Modal Attention, logic ngắt SOS, chủ biên báo cáo.")
    add_bullet("Nguyễn Duy Khang: ", "Computer Vision Lead, tiền xử lý ảnh nón bảo hiểm/hung khí (NTHU-DDD, Roboflow), MobileNetV3-Small backbone.")
    add_bullet("Phạm Thành Trung: ", "Bio-Signals Specialist, tiền xử lý chuỗi PPG (WESAD, Smart Band), trích xuất đặc trưng HRV trong hrv.py, huấn luyện BiLSTM.")
    add_bullet("Vũ Văn Duy: ", "ML Evaluation Lead, tối ưu hóa Focal Loss, Class-Balanced Loss, Cost-Sensitive Bayes Matrix, kiểm định Stratified 5-Fold.")
    add_bullet("Hoàng Cao Sơn: ", "Data Engineering & XAI, trực quan hóa Grad-CAM và SHAP, thực nghiệm hiện trường TP.HCM, dựng video clip demo.")

    add_h2("3. Thông qua Kế hoạch Thực hiện 8 Tuần (Tuần 2 đến Tuần 9)")
    add_p("Nhóm nhất trí 100% với lộ trình 5 giai đoạn: Tuần 2-3 (Dữ liệu & Augmentation), Tuần 4 (Baselines & HRV), Tuần 5-6 (Deep Multimodal & Focal Loss), Tuần 7 (Thẩm định 5-Fold, TP.HCM & XAI), Tuần 8-9 (Hoàn thiện báo cáo >= 40 trang & Bảo vệ).")

    add_h2("4. Thống nhất Kỷ luật Git & Quy chuẩn Đóng góp Mã nguồn")
    add_bullet("Kho lưu trữ GitHub: ", "https://github.com/mrkiss-it/multimodal-driver-safety-ml")
    add_bullet("Quy tắc nhánh: ", "Không push trực tiếp vào main. Mọi tính năng phát triển trên nhánh feature/* rẽ nhánh từ develop. Merge qua Pull Request có review.")
    add_bullet("Quy tắc Commit: ", "Tuân thủ Conventional Commits (feat:, fix:, docs:, test:, refactor:).")
    add_bullet("Kiểm soát dữ liệu: ", "Nghiêm cấm commit file dữ liệu thô nặng (>50MB) vào Git.")

    add_h2("5. Đầu việc Cụ thể trong Tuần 2 (27/10 – 01/11/2026)")
    add_bullet("Hôn: ", "Hoàn thiện CONTRIBUTING.md, PR/Issue Templates, nhánh develop và Biên bản họp khởi động.")
    add_bullet("Khang: ", "Tải NTHU-DDD và Roboflow Weapon về máy, kiểm tra số lượng và định dạng ảnh.")
    add_bullet("Trung: ", "Tải WESAD, viết script đọc file .pkl trích xuất chuỗi nhịp tim PPG.")
    add_bullet("Duy: ", "Tạo requirements.txt, viết test_env.py kiểm tra CUDA GPU và thư viện PyTorch/Scikit-learn.")
    add_bullet("Sơn: ", "Lên kế hoạch thu thập ảnh nón bảo hiểm ngoài đường phố TP.HCM, khởi tạo notebook EDA.")

    # III. KẾT LUẬN & BIỂU QUYẾT
    add_h1("III. KẾT LUẬN & BIỂU QUYẾT")
    add_p(
        "Toàn thể 05/05 thành viên nhóm tham gia cuộc họp đã nghe, thảo luận dân chủ và nhất trí biểu quyết thông qua 100% "
        "các nội dung trong biên bản này. Biên bản được lập thành văn bản lưu vào thư mục docs/ và công khai trên GitHub repo của nhóm."
    )
    add_p("Cuộc họp kết thúc tốt đẹp vào lúc 21h45 cùng ngày.")

    # Chữ ký
    p_sign_t = doc.add_paragraph()
    p_sign_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_date = p_sign_t.add_run("TP. Hồ Chí Minh, ngày 27 tháng 10 năm 2026\n")
    r_date.font.italic = True
    r_date.font.size = Pt(10)
    p_sign_t.paragraph_format.space_before = Pt(12)
    p_sign_t.paragraph_format.space_after = Pt(6)

    sign_table = doc.add_table(rows=3, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = sign_table.cell(0, 0)
    c1 = sign_table.cell(0, 1)
    
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s0 = p0.add_run("THƯ KÝ CUỘC HỌP\n(Ký và ghi rõ họ tên)\n\n\n\n")
    r_s0.font.bold = True
    r_s0.font.size = Pt(10)
    r_name0 = p0.add_run("Hoàng Cao Sơn")
    r_name0.font.bold = True
    r_name0.font.size = Pt(10.5)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s1 = p1.add_run("CHỦ TỌA / TRƯỞNG NHÓM\n(Ký và ghi rõ họ tên)\n\n\n\n")
    r_s1.font.bold = True
    r_s1.font.size = Pt(10)
    r_name1 = p1.add_run("Trần Hoàng Hôn")
    r_name1.font.bold = True
    r_name1.font.size = Pt(10.5)

    set_table_borders(sign_table, "FFFFFF")

    docs_output = r"D:\Trí tuệ nhân tạo\HK3\Máy học nâng cao\Đồ án\docs\Bien_Ban_Hop_Khoi_Dong_Tuan_2.docx"
    doc.save(docs_output)
    print("Meeting minutes docx created successfully.")

if __name__ == "__main__":
    create_meeting_document()
